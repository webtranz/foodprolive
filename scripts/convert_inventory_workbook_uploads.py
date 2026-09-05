import csv
import json
import re
import sys
from pathlib import Path

import openpyxl


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_INPUT = Path(r"C:\Users\HP\Downloads\kbr-384 31-8-29 inventory.xlsx")
OUT_DIR = ROOT / "artifacts" / "inventory-ingredients-only" / "kbr-384-31-8-29"
INGREDIENTS_OUTPUT = OUT_DIR / "kbr-384-31-8-29-ingredients-upload.csv"
INVENTORY_OUTPUT = OUT_DIR / "kbr-384-31-8-29-inventory-upload.csv"
QC_OUTPUT = OUT_DIR / "kbr-384-31-8-29-conversion-qc.csv"

SITE = "KBR"
WAREHOUSE = "KBR-384"

INVENTORY_HEADERS = [
    "site",
    "warehouse",
    "item_code",
    "ingredient_name",
    "quantity",
    "unit",
    "unit_cost",
    "reorder_level",
    "status",
    "item_group",
    "amount",
]

INGREDIENT_HEADERS = [
    "item_group",
    "ingredient_code",
    "sku",
    "ingredient_name",
    "unit",
    "unit_price",
    "cooking_yield_percent",
    "calories_per_100g",
    "allergens",
    "package_base_quantity",
    "package_base_unit",
    "package_parse_source",
]

HEADER_ALIASES = {
    "item_group": {"item_group", "group", "category", "item group"},
    "item_code": {"item", "item_code", "item code", "ingredient_code", "ingredient code", "sku"},
    "item_duplicate": {"item_duplicate", "item duplicate"},
    "ingredient_name": {"product_name", "product name", "ingredient_name", "ingredient name", "name", "description"},
    "unit": {"unit", "uom"},
    "site": {"site", "project"},
    "warehouse": {"warehouse", "store", "location"},
    "unit_cost": {"unit_price", "unit price", "price", "unit_cost", "unit cost", "cost"},
    "quantity": {"qty", "quantity", "on hand", "on_hand", "stock", "stock on hand"},
    "quantity_duplicate": {"qty_duplicate", "qty duplicate"},
    "amount": {"amount", "value", "total", "total value"},
}

UNIT_ALIASES = {
    "kgs": "KG",
    "kg": "KG",
    "kilogram": "KG",
    "kilograms": "KG",
    "g": "G",
    "gram": "G",
    "grams": "G",
    "ltr": "LTR",
    "lt": "LTR",
    "l": "LTR",
    "liter": "LTR",
    "liters": "LTR",
    "litre": "LTR",
    "litres": "LTR",
    "ml": "ML",
    "ea": "EA",
    "each": "EA",
    "pak": "PAK",
    "pack": "PAK",
    "packet": "PAK",
    "pkt": "PAK",
    "cs": "CS",
    "case": "CS",
    "bdl": "BDL",
    "bundle": "BDL",
}


def clean_header(value):
    return re.sub(r"[\s_\-]+", " ", str(value or "").strip().lower())


def clean_text(value):
    return re.sub(r"\s+", " ", str(value or "").strip())


def parse_number(value):
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).replace(",", "").strip()
    match = re.search(r"-?\d+(?:\.\d+)?", text)
    return float(match.group(0)) if match else None


def format_number(value):
    if value is None:
        return ""
    value = round(float(value), 6)
    text = f"{value:.6f}".rstrip("0").rstrip(".")
    return text or "0"


def normalize_unit(value):
    text = clean_text(value)
    return UNIT_ALIASES.get(text.lower(), text.upper())


def find_header_row(ws):
    best = None
    best_score = -1
    for row_index, row in enumerate(ws.iter_rows(min_row=1, max_row=min(ws.max_row, 25), values_only=True), 1):
        headers = [clean_header(cell) for cell in row]
        score = 0
        for aliases in HEADER_ALIASES.values():
            if any(header in aliases for header in headers):
                score += 1
        if score > best_score:
            best = row_index
            best_score = score
    return best if best_score >= 3 else None


def header_map(headers):
    mapped = {}
    for index, header in enumerate(headers):
        clean = clean_header(header)
        for field, aliases in HEADER_ALIASES.items():
            if clean in aliases and field not in mapped:
                mapped[field] = index
    return mapped


def cell(row, mapping, field):
    index = mapping.get(field)
    if index is None or index >= len(row):
        return None
    return row[index]


def parse_package_fields(name, unit):
    unit = normalize_unit(unit)
    if unit == "BDL":
        return {
            "package_base_quantity": "0.08",
            "package_base_unit": "kg",
            "package_parse_source": "default_bundle_weight",
        }
    text = str(name or "").upper()
    numeric = r"\d+(?:\.\d+)?(?:\s*-\s*\d+(?:\.\d+)?)?"
    measure = r"KG|KGS|G|GM|GMS|GRAMS|LTR|L|LT|LITRE|LITER|ML|CT|CNT|COUNT|COUNTS|OZ|Z"
    tail = re.search(rf"((?:{numeric}\s*/\s*)+)({numeric})\s*({measure})\b", text)
    if not tail:
        return {}
    counts = [
        float(part.strip().split("-")[0])
        for part in tail.group(1).split("/")
        if part.strip()
    ]
    size = float(tail.group(2).strip().split("-")[0])
    size_unit = tail.group(3).lower()
    multiplier = 1
    lower_unit = unit.lower()
    if lower_unit in {"cs", "case"}:
        for count in counts:
            multiplier *= count
    elif lower_unit in {"pak", "pkt", "pack", "packet"} and len(counts) > 1:
        for count in counts[1:]:
            multiplier *= count
    base_quantity = size * multiplier
    if size_unit in {"kg", "kgs"}:
        base_quantity = base_quantity
        base_unit = "kg"
    elif size_unit in {"g", "gm", "gms", "grams"}:
        base_quantity = base_quantity / 1000
        base_unit = "kg"
    elif size_unit in {"ltr", "l", "lt", "litre", "liter"}:
        base_unit = "l"
    elif size_unit == "ml":
        base_quantity = base_quantity / 1000
        base_unit = "l"
    elif size_unit in {"ct", "cnt", "count", "counts"}:
        base_unit = "pieces"
    elif size_unit in {"oz", "z"}:
        base_quantity = base_quantity * 0.028349523125
        base_unit = "kg"
    else:
        return {}
    return {
        "package_base_quantity": format_number(base_quantity),
        "package_base_unit": base_unit,
        "package_parse_source": "item_name_package",
    }


def load_previous_ingredient_defaults():
    defaults = {}
    for path in [
        ROOT / "artifacts" / "inventory-ingredients-only" / "ingredients-yield-calories-allergens-bdl-80g.csv",
        ROOT / "artifacts" / "inventory-ingredients-only" / "kbr-384-inventory-upload-corrected.csv",
    ]:
        if not path.exists():
            continue
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            reader = csv.DictReader(handle)
            for row in reader:
                code = clean_text(row.get("ingredient_code") or row.get("item_code") or row.get("sku"))
                if code and code not in defaults:
                    defaults[code] = row
    return defaults


def workbook_rows(path):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    records = []
    sheet_info = []
    for ws in wb.worksheets:
        header_row = find_header_row(ws)
        sheet_info.append({"sheet": ws.title, "rows": ws.max_row, "cols": ws.max_column, "header_row": header_row})
        if not header_row:
            continue
        rows = list(ws.iter_rows(min_row=header_row, values_only=True))
        if not rows:
            continue
        headers = rows[0]
        mapping = header_map(headers)
        for excel_row, row in enumerate(rows[1:], header_row + 1):
            item_code = clean_text(cell(row, mapping, "item_code"))
            name = clean_text(cell(row, mapping, "ingredient_name"))
            quantity = parse_number(cell(row, mapping, "quantity"))
            unit = normalize_unit(cell(row, mapping, "unit"))
            if not item_code or not name or quantity is None or not unit:
                continue
            unit_cost = parse_number(cell(row, mapping, "unit_cost"))
            amount = parse_number(cell(row, mapping, "amount"))
            records.append({
                "source_sheet": ws.title,
                "source_row": excel_row,
                "item_group": clean_text(cell(row, mapping, "item_group")),
                "item_code": item_code,
                "ingredient_name": name,
                "quantity": quantity,
                "unit": unit,
                "unit_cost": unit_cost,
                "amount": amount,
                "site": clean_text(cell(row, mapping, "site")) or SITE,
                "warehouse": clean_text(cell(row, mapping, "warehouse")) or WAREHOUSE,
            })
    return records, sheet_info


def main():
    input_path = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_INPUT
    defaults = load_previous_ingredient_defaults()
    rows, sheet_info = workbook_rows(input_path)
    if not rows:
        raise SystemExit(f"No inventory rows found in {input_path}")

    inventory_rows = []
    ingredient_rows = []
    qc_rows = []
    seen_ingredients = set()
    seen_inventory = set()

    for source in rows:
        code = source["item_code"]
        default = defaults.get(code, {})
        quantity = source["quantity"]
        unit_cost = source["unit_cost"]
        amount = source["amount"] if source["amount"] is not None else (
            quantity * unit_cost if unit_cost is not None else None
        )
        status = "in_stock" if quantity > 0 else "out_of_stock"
        inventory_key = (source["warehouse"], code)
        if inventory_key not in seen_inventory:
            seen_inventory.add(inventory_key)
            inventory_rows.append({
                "site": source["site"] or SITE,
                "warehouse": source["warehouse"] or WAREHOUSE,
                "item_code": code,
                "ingredient_name": source["ingredient_name"],
                "quantity": format_number(quantity),
                "unit": source["unit"],
                "unit_cost": format_number(unit_cost),
                "reorder_level": "0",
                "status": status,
                "item_group": source["item_group"] or default.get("category") or default.get("item_group") or "",
                "amount": format_number(amount),
            })
        if code not in seen_ingredients:
            seen_ingredients.add(code)
            package = parse_package_fields(source["ingredient_name"], source["unit"])
            ingredient_rows.append({
                "item_group": source["item_group"] or default.get("category") or default.get("item_group") or "",
                "ingredient_code": code,
                "sku": code,
                "ingredient_name": source["ingredient_name"],
                "unit": source["unit"],
                "unit_price": format_number(unit_cost),
                "cooking_yield_percent": default.get("cooking_yield_percent") or "100",
                "calories_per_100g": default.get("calories_per_100g") or "",
                "allergens": default.get("allergens") or "none",
                "package_base_quantity": package.get("package_base_quantity", default.get("package_base_quantity", "")),
                "package_base_unit": package.get("package_base_unit", default.get("package_base_unit", "")),
                "package_parse_source": package.get("package_parse_source", default.get("package_parse_source", "")),
            })
        qc_rows.append({
            "source_sheet": source["source_sheet"],
            "source_row": source["source_row"],
            "item_code": code,
            "ingredient_name": source["ingredient_name"],
            "quantity": format_number(quantity),
            "unit": source["unit"],
            "unit_cost": format_number(unit_cost),
            "amount": format_number(amount),
            "status": "converted",
        })

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    with INVENTORY_OUTPUT.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=INVENTORY_HEADERS)
        writer.writeheader()
        writer.writerows(inventory_rows)
    with INGREDIENTS_OUTPUT.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=INGREDIENT_HEADERS)
        writer.writeheader()
        writer.writerows(ingredient_rows)
    with QC_OUTPUT.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=[
            "source_sheet", "source_row", "item_code", "ingredient_name",
            "quantity", "unit", "unit_cost", "amount", "status",
        ])
        writer.writeheader()
        writer.writerows(qc_rows)

    print(json.dumps({
        "input": str(input_path),
        "sheet_info": sheet_info,
        "source_rows": len(rows),
        "inventory_rows": len(inventory_rows),
        "ingredient_rows": len(ingredient_rows),
        "inventory_output": str(INVENTORY_OUTPUT),
        "ingredients_output": str(INGREDIENTS_OUTPUT),
        "qc_output": str(QC_OUTPUT),
    }, indent=2))


if __name__ == "__main__":
    main()
